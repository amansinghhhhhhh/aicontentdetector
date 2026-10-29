"""Extract plain text from uploaded files (PDF / DOCX / TXT / MD)."""

import io
from typing import Set

SUPPORTED_EXTENSIONS: Set[str] = {".txt", ".md", ".pdf", ".docx"}


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    if not filename:
        raise UnsupportedFileType("Missing filename.")
    name = filename.lower()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""

    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or name}'. Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext in (".txt", ".md"):
        text = _decode(data)
    elif ext == ".pdf":
        text = _extract_pdf(data)
    else:
        text = _extract_docx(data)

    text = text.strip()
    if not text:
        raise EmptyDocument(f"No extractable text found in '{filename}'.")
    return text
