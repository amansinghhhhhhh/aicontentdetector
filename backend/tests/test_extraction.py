import io

import pytest

from app.services.extraction import (
    EmptyDocument,
    UnsupportedFileType,
    extract_text,
)


def test_txt_utf8():
    assert extract_text("a.txt", "hello café".encode("utf-8")) == "hello café"


def test_txt_cp1252_fallback():
    assert extract_text("a.txt", "price: 5€".encode("cp1252")) == "price: 5€"


def test_md_supported():
    assert extract_text("README.md", b"# Title\n\nBody") == "# Title\n\nBody"


def test_unsupported_extension():
    with pytest.raises(UnsupportedFileType):
        extract_text("doc.rtf", b"data")


def test_missing_filename():
    with pytest.raises(UnsupportedFileType):
        extract_text("", b"data")


def test_empty_document():
    with pytest.raises(EmptyDocument):
        extract_text("empty.txt", b"   \n\t ")


def test_docx_extraction():
    from docx import Document

    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("report.docx", buffer.getvalue())
    assert "First paragraph." in text
    assert "Second paragraph." in text
